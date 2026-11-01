"""A provider's financial report as an Excel workbook: summary, earnings, withdrawals and a guide.

Every figure comes from the same ledger as the earnings screen, for one period in Addis Ababa time.
Text that people typed (job titles, client names, notes) is written as plain text, never as a formula.
"""

from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.db import fetch_all, fetch_one

from .earnings import LEDGER_SELECT, PAYOUT_COLUMNS, _balances, _test_money
from .gateways import PaymentError

TZ = ZoneInfo("Africa/Addis_Ababa")
MAX_DAYS = 366
MONEY = '#,##0.00 "ETB";-#,##0.00 "ETB"'
GREEN = "0B6B3A"
HEAD_FILL = PatternFill("solid", fgColor="E6F2EA")

KIND_LABEL = {
    "earning": "Earning from a paid job",
    "payout": "Withdrawal to your bank",
    "payout_return": "Withdrawal that came back",
    "refund": "Refund to a client",
    "adjustment": "Correction by our team",
}
PAYOUT_STATUS = {
    "requested": "Waiting for approval", "approved": "On its way to your bank", "paid": "Paid to your bank",
    "rejected": "Not approved", "failed": "Failed, money back in your balance",
    "returned": "Came back from the bank, money back in your balance",
}


def _text(value) -> str:
    """Plain text for a cell: anything a spreadsheet could read as a formula gets a leading apostrophe."""
    s = "" if value is None else str(value)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


def _local(ts: datetime | None) -> datetime | None:
    # Excel has no time zones: show Addis Ababa wall-clock time without the offset
    return ts.astimezone(TZ).replace(tzinfo=None) if ts else None


def _period(date_from: date | None, date_to: date | None) -> tuple[date, date]:
    today = datetime.now(TZ).date()
    end = date_to or today
    start = date_from or date(end.year, 1, 1)
    if start > end:
        raise PaymentError("The start date must be on or before the end date")
    if (end - start).days >= MAX_DAYS:
        raise PaymentError("Choose a period of at most one year")
    return start, end


def _bounds(start: date, end: date) -> tuple[datetime, datetime]:
    return datetime.combine(start, time.min, TZ), datetime.combine(end + timedelta(days=1), time.min, TZ)


def build(user_id: str, date_from: date | None = None, date_to: date | None = None) -> tuple[bytes, str]:
    """Returns (xlsx bytes, file name)."""
    from services import provider_service

    provider = provider_service.get_for_user(user_id)
    if not provider:
        raise PaymentError("Join as a provider first", 404)
    start, end = _period(date_from, date_to)
    lo, hi = _bounds(start, end)
    pid = provider["id"]
    # Same rule as the balance: once live, test money is left out of the figures
    only = "" if _test_money() else "and not l.is_test"

    opening = fetch_one(f"select coalesce(sum(l.amount), 0)::float as v from provider_ledger l "
                        f"where l.provider_id = %s and l.created_at < %s {only}", (pid, lo))["v"]
    lines = fetch_all(f"{LEDGER_SELECT} where l.provider_id = %s and l.created_at >= %s and l.created_at < %s {only} "
                      f"order by l.created_at, l.id", (pid, lo, hi))
    payouts = fetch_all(f"select {PAYOUT_COLUMNS} from payouts po where po.provider_id = %s "
                        f"and po.created_at >= %s and po.created_at < %s "
                        f"{'' if _test_money() else 'and not po.is_test'} order by po.created_at", (pid, lo, hi))
    now = _balances(pid)

    def total(kind: str) -> float:
        return round(sum(r["amount"] for r in lines if r["kind"] == kind), 2)

    earned = [r for r in lines if r["kind"] == "earning"]
    closing = round(opening + sum(r["amount"] for r in lines), 2)

    wb = Workbook()
    _summary(wb.active, provider, start, end, {
        "opening": opening,
        "client_paid": sum((r["gross"] or 0) + (r["provider_vat"] or 0) for r in earned),
        "price": sum(r["gross"] or 0 for r in earned),
        "provider_vat": sum(r["provider_vat"] or 0 for r in earned),
        "commission": sum(r["commission"] or 0 for r in earned),
        "commission_vat": sum(r["commission_vat"] or 0 for r in earned),
        "earned": total("earning"), "refunds": total("refund") + total("adjustment"),
        "payouts": total("payout"), "returns": total("payout_return"), "closing": closing,
        "jobs": len(earned), "now": now,
    })
    _earnings_sheet(wb.create_sheet("Earnings"), lines)
    _payouts_sheet(wb.create_sheet("Withdrawals"), payouts)
    _guide(wb.create_sheet("How to read this"))

    buf = io.BytesIO()
    wb.save(buf)
    slug = "".join(c for c in (provider.get("slug") or "provider") if c.isalnum() or c == "-")[:40] or "provider"
    return buf.getvalue(), f"earnings-{slug}-{start:%Y-%m-%d}-to-{end:%Y-%m-%d}.xlsx"


def _header(ws, row: int, titles: list[str]) -> None:
    for col, title in enumerate(titles, 1):
        c = ws.cell(row=row, column=col, value=title)
        c.font, c.fill = Font(bold=True), HEAD_FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _summary(ws, provider: dict, start: date, end: date, f: dict) -> None:
    ws.title = "Summary"
    ws["A1"] = "Financial report · የገንዘብ ሪፖርት"
    ws["A1"].font = Font(bold=True, size=14, color=GREEN)
    ws["A2"] = _text(provider.get("display_name") or "")
    ws["A3"] = f"Period: {start:%d %b %Y} to {end:%d %b %Y} (Addis Ababa time)"
    ws["A4"] = f"Made on {datetime.now(TZ):%d %b %Y %H:%M}"
    rows = [
        ("Balance at the start of the period", f["opening"], "What you had before the first day"),
        ("Paid jobs", f["jobs"], None),
        ("Clients paid in total", f["client_paid"], "Your prices plus the VAT you charge, if you are VAT-registered"),
        ("Your prices", f["price"], None),
        ("VAT you collected from clients", f["provider_vat"], "Yours to declare to the tax office"),
        ("Platform commission", -f["commission"], "Our fee, a percentage of your price"),
        ("VAT on the commission", -f["commission_vat"], "15% VAT on our fee, not on your price"),
        ("Your earnings", f["earned"], "Price minus commission and its VAT, plus the VAT you collected"),
        ("Refunds to clients", f["refunds"], "Taken back from your balance"),
        ("Withdrawals paid to your bank", f["payouts"], None),
        ("Withdrawals that came back", f["returns"], "Back in your balance after the bank returned them"),
        ("Balance at the end of the period", f["closing"], "Start balance plus everything above"),
        (None, None, None),
        ("Today: on hold (dispute window)", f["now"]["held"], "Released 7 days after the client paid"),
        ("Today: waiting for payout", f["now"]["in_payout"], None),
        ("Today: available to withdraw", max(f["now"]["available"], 0), None),
    ]
    _header(ws, 6, ["Item", "Amount", "What it means"])
    for i, (label, value, note) in enumerate(rows, 7):
        if label is None:
            continue
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=round(value, 2) if isinstance(value, float) else value)
        if label != "Paid jobs":
            c.number_format = MONEY
        if note:
            ws.cell(row=i, column=3, value=note)
        if label.startswith(("Your earnings", "Balance at the end")):
            ws.cell(row=i, column=1).font = ws.cell(row=i, column=2).font = Font(bold=True)
    ws.freeze_panes = None
    _widths(ws, [38, 20, 62])


def _earnings_sheet(ws, lines: list[dict]) -> None:
    _header(ws, 1, ["Date", "What", "Job", "Client", "Your price", "VAT you charged", "Client paid",
                    "Commission", "VAT on commission", "Amount to you", "Reference", "Test"])
    for r, line in enumerate(lines, 2):
        price, pvat = line.get("gross"), line.get("provider_vat")
        values = [
            _local(line["created_at"]), KIND_LABEL.get(line["kind"], line["kind"]), _text(line.get("job_title")),
            _text((line.get("client_name") or "").split(" ")[0] or None),
            price, pvat, (price or 0) + (pvat or 0) if price is not None else None,
            -(line["commission"] or 0) if line.get("commission") is not None else None,
            -(line["commission_vat"] or 0) if line.get("commission_vat") is not None else None,
            line["amount"],
            _text(line.get("tx_ref") or line.get("payout_reference") or line.get("refund_reference")),
            "Test" if line.get("is_test") else "",
        ]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=r, column=col, value=v)
            if col == 1 and v:
                c.number_format = "dd mmm yyyy hh:mm"
            elif 5 <= col <= 10 and v is not None:
                c.number_format = MONEY
    if not lines:
        ws.cell(row=2, column=1, value="No money moved in this period.")
    _widths(ws, [18, 26, 36, 14, 15, 15, 15, 15, 17, 16, 26, 7])


def _payouts_sheet(ws, payouts: list[dict]) -> None:
    _header(ws, 1, ["Asked on", "Amount", "Status", "Bank", "Account", "Reference", "Paid on", "Note", "Test"])
    for r, p in enumerate(payouts, 2):
        values = [_local(p["created_at"]), p["amount"], PAYOUT_STATUS.get(p["status"], p["status"]),
                  _text(p.get("bank_name")), f"•••• {p.get('account_last4') or ''}", _text(p.get("reference")),
                  _local(p.get("paid_at")), _text(p.get("note")), "Test" if p.get("is_test") else ""]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=r, column=col, value=v)
            if col in (1, 7) and v:
                c.number_format = "dd mmm yyyy hh:mm"
            elif col == 2:
                c.number_format = MONEY
    if not payouts:
        ws.cell(row=2, column=1, value="No withdrawals asked for in this period.")
    _widths(ws, [18, 16, 40, 28, 12, 22, 18, 36, 7])


def _guide(ws) -> None:
    lines = [
        ("How to read this report", "ይህን ሪፖርት እንዴት እንደሚያነቡት"),
        ("Your price is what you and the client agreed. If you are VAT-registered, the client also pays 15% VAT on it, "
         "which is yours to declare.", "ዋጋዎ እርስዎና ደንበኛው የተስማማችሁበት ነው።"),
        ("Platform commission is our fee, a percentage of your price. VAT of 15% is charged on the commission only.",
         "የመድረኩ ኮሚሽን ከዋጋዎ የሚወሰድ ክፍያ ነው፤ ተ.እ.ታ የሚከፈለው በኮሚሽኑ ላይ ብቻ ነው።"),
        ("Amount to you is what was added to (or taken from) your balance by that line.", ""),
        ("New earnings are on hold for 7 days after the client pays (the dispute window), then you can withdraw them.",
         "አዲስ ገቢ ደንበኛው ከከፈለ ከ7 ቀን በኋላ ለማውጣት ዝግጁ ይሆናል።"),
        ("Withdrawals have no fee and no VAT: it is your own money moving to your bank.",
         "ገንዘብ ማውጣት ክፍያም ተ.እ.ታም የለውም።"),
        ("Times are Addis Ababa time. Lines marked Test are from test payments and never count as real money.", ""),
        ("Questions: open Help in your account, or reply to any email from us.", ""),
    ]
    for r, (en, am) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=en).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=2, value=am).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=13, color=GREEN)
    _widths(ws, [80, 60])
