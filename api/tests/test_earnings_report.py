"""The provider's Excel financial report: figures match the ledger, typed text can't become a formula,
only the owner can download it, and the period rules hold."""

import io
from datetime import date, timedelta

from openpyxl import load_workbook

from conftest import API
from services.db import execute
from test_money_safety import _paid_job

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _download(client, user, **params):
    return client.get(f"{API}/providers/me/earnings/report.xlsx", params=params, headers=user["headers"])


def _summary(wb) -> dict:
    ws = wb["Summary"]
    return {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(7, ws.max_row + 1)}


def test_report_matches_the_ledger(client, make_user):
    job = _paid_job(client, make_user, price=1000)  # 1000 price, 100 commission, 15 VAT on it: 885 earned
    r = _download(client, job["agent"])
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX)
    assert r.headers["content-disposition"].startswith('attachment; filename="earnings-')
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Earnings", "Withdrawals", "How to read this"]
    s = _summary(wb)
    assert s["Paid jobs"] == 1
    assert s["Your prices"] == 1000
    assert s["Platform commission"] == -100
    assert s["Your earnings"] == 885
    assert s["Balance at the end of the period"] == 885
    row = [c.value for c in wb["Earnings"][2]]
    assert row[1] == "Earning from a paid job" and row[4] == 1000 and row[7] == -100 and row[9] == 885


def test_typed_text_is_never_a_formula(client, make_user):
    job = _paid_job(client, make_user, title='=HYPERLINK("http://evil.example","click")')
    wb = load_workbook(io.BytesIO(_download(client, job["agent"]).content))
    job_cell = wb["Earnings"]["C2"]
    assert job_cell.data_type == "s" and job_cell.value.startswith("'=")


def test_only_providers_get_a_report_and_the_period_is_checked(client, make_user):
    stranger = make_user()
    assert _download(client, stranger).status_code == 404  # not a provider: nothing to show
    job = _paid_job(client, make_user)
    today = date.today()
    assert _download(client, job["agent"], **{"from": str(today), "to": str(today - timedelta(days=1))}).status_code == 400
    assert _download(client, job["agent"], **{"from": str(today - timedelta(days=400)), "to": str(today)}).status_code == 400
    # A period before the job: nothing moved, the start and end balances are 0
    old = _download(client, job["agent"], **{"from": "2020-01-01", "to": "2020-12-31"})
    s = _summary(load_workbook(io.BytesIO(old.content)))
    assert s["Paid jobs"] == 0 and s["Balance at the end of the period"] == 0


def test_opening_balance_carries_earlier_money(client, make_user):
    job = _paid_job(client, make_user)
    execute("update provider_ledger set created_at = now() - interval '40 days' where provider_id = %s",
            (job["provider"]["id"],))
    today = date.today()
    s = _summary(load_workbook(io.BytesIO(_download(
        client, job["agent"], **{"from": str(today - timedelta(days=10)), "to": str(today)}).content)))
    assert s["Balance at the start of the period"] == 885 and s["Paid jobs"] == 0
    assert s["Balance at the end of the period"] == 885
